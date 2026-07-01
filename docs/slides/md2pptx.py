#!/usr/bin/env python3
"""Convert markdown slides to PowerPoint."""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
import re
import os

def parse_slides(md_content):
    """Parse markdown content into slides."""
    slides = []
    
    # Split by horizontal rule (slide separator)
    parts = re.split(r'\n---\n', md_content)
    
    for i, part in enumerate(parts):
        lines = part.strip().split('\n')
        if not lines:
            continue
        
        # Extract slide number and title from first lines
        slide = {
            'number': i + 1,
            'title': '',
            'content': []
        }
        
        # Skip header if exists
        start_idx = 0
        if lines[0].startswith('# '):
            slide['title'] = lines[0][2:].strip()
            start_idx = 1
        
        # Parse remaining content
        current_section = {'title': '', 'items': []}
        in_code_block = False
        code_content = []
        
        for line in lines[start_idx:]:
            # Toggle code block
            if line.strip().startswith('```'):
                if in_code_block:
                    # End code block
                    current_section['items'].append({
                        'type': 'code',
                        'content': '\n'.join(code_content)
                    })
                    code_content = []
                else:
                    # Start code block
                    if current_section['items']:
                        if current_section['title'] or current_section['items']:
                            slide['content'].append(current_section)
                    current_section = {'title': '', 'items': []}
                in_code_block = not in_code_block
                continue
            
            if in_code_block:
                code_content.append(line)
                continue
            
            # Headers
            if line.startswith('## '):
                if current_section['title'] or current_section['items']:
                    slide['content'].append(current_section)
                current_section = {'title': line[3:].strip(), 'items': []}
            elif line.startswith('### '):
                current_section['items'].append({
                    'type': 'subsection',
                    'content': line[4:].strip()
                })
            elif line.startswith('- '):
                current_section['items'].append({
                    'type': 'bullet',
                    'content': line[2:].strip()
                })
            elif line.startswith('| '):
                # Table - collect all lines
                table_lines = [line]
                j = lines.index(line) + 1
                while j < len(lines) and lines[j].startswith('|'):
                    table_lines.append(lines[j])
                    j += 1
                current_section['items'].append({
                    'type': 'table',
                    'content': table_lines
                })
            elif line.strip():
                current_section['items'].append({
                    'type': 'text',
                    'content': line.strip()
                })
        
        if current_section['title'] or current_section['items']:
            slide['content'].append(current_section)
        
        slides.append(slide)
    
    return slides

def create_presentation(slides, output_path):
    """Create PowerPoint presentation from slides."""
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    
    # Colors
    PRIMARY = RGBColor(21, 101, 192)  # #1565C0
    DARK = RGBColor(27, 42, 74)       # #1B2A4A
    WHITE = RGBColor(255, 255, 255)
    GRAY = RGBColor(100, 116, 139)   # #64748B
    
    for slide_data in slides:
        # Add blank slide
        slide_layout = prs.slide_layouts[6]  # Blank
        slide = prs.slides.add_slide(slide_layout)
        
        # Add header background
        header = slide.shapes.add_shape(
            1,  # Rectangle
            Inches(0), Inches(0),
            Inches(13.333), Inches(2)
        )
        header.fill.solid()
        header.fill.fore_color.rgb = PRIMARY
        header.line.fill.background()
        
        # Add slide number
        num_box = slide.shapes.add_textbox(
            Inches(0.5), Inches(0.3),
            Inches(1), Inches(0.4)
        )
        num_frame = num_box.text_frame
        num_para = num_frame.paragraphs[0]
        num_para.text = f"Slide {slide_data['number']}"
        num_para.font.size = Pt(14)
        num_para.font.color.rgb = WHITE
        num_para.font.bold = True
        
        # Add title
        title_box = slide.shapes.add_textbox(
            Inches(0.5), Inches(0.7),
            Inches(12), Inches(1)
        )
        title_frame = title_box.text_frame
        title_para = title_frame.paragraphs[0]
        title_para.text = slide_data['title']
        title_para.font.size = Pt(32)
        title_para.font.color.rgb = WHITE
        title_para.font.bold = True
        
        # Add content area
        content_top = Inches(2.2)
        content_height = Inches(5)
        
        y_pos = content_top.inches
        
        for section in slide_data['content']:
            # Section title
            if section['title']:
                sec_title = slide.shapes.add_textbox(
                    Inches(0.5), Inches(y_pos),
                    Inches(12), Inches(0.4)
                )
                sec_frame = sec_title.text_frame
                sec_para = sec_frame.paragraphs[0]
                sec_para.text = section['title']
                sec_para.font.size = Pt(18)
                sec_para.font.color.rgb = PRIMARY
                sec_para.font.bold = True
                y_pos += 0.4
            
            # Section items
            for item in section['items']:
                if y_pos > 6.5:
                    break
                
                if item['type'] == 'bullet':
                    bullet = slide.shapes.add_textbox(
                        Inches(0.8), Inches(y_pos),
                        Inches(11.5), Inches(0.35)
                    )
                    bullet_frame = bullet.text_frame
                    bullet_para = bullet_frame.paragraphs[0]
                    bullet_para.text = f"• {item['content']}"
                    bullet_para.font.size = Pt(14)
                    bullet_para.font.color.rgb = DARK
                    y_pos += 0.3
                
                elif item['type'] == 'text':
                    text = slide.shapes.add_textbox(
                        Inches(0.8), Inches(y_pos),
                        Inches(11.5), Inches(0.35)
                    )
                    text_frame = text.text_frame
                    text_para = text_frame.paragraphs[0]
                    text_para.text = item['content']
                    text_para.font.size = Pt(14)
                    text_para.font.color.rgb = DARK
                    y_pos += 0.3
                
                elif item['type'] == 'code':
                    # Code block - simplified
                    code = slide.shapes.add_textbox(
                        Inches(0.5), Inches(y_pos),
                        Inches(12), Inches(1.5)
                    )
                    code_frame = code.text_frame
                    code_frame.word_wrap = True
                    code_para = code_frame.paragraphs[0]
                    # Truncate long code
                    code_text = item['content'][:500] + ('...' if len(item['content']) > 500 else '')
                    code_para.text = code_text
                    code_para.font.size = Pt(9)
                    code_para.font.name = 'Courier New'
                    code_para.font.color.rgb = RGBColor(100, 100, 100)
                    
                    # Add background
                    bg = slide.shapes.add_shape(
                        1, Inches(0.5), Inches(y_pos - 0.05),
                        Inches(12), Inches(1.6)
                    )
                    bg.fill.solid()
                    bg.fill.fore_color.rgb = RGBColor(30, 30, 30)
                    bg.line.fill.background()
                    # Move to back
                    spTree = slide.shapes._spTree
                    sp = bg._element
                    spTree.remove(sp)
                    spTree.insert(2, sp)
                    
                    y_pos += 1.6
                
                elif item['type'] == 'subsection':
                    sub = slide.shapes.add_textbox(
                        Inches(0.8), Inches(y_pos),
                        Inches(11), Inches(0.35)
                    )
                    sub_frame = sub.text_frame
                    sub_para = sub_frame.paragraphs[0]
                    sub_para.text = f"▸ {item['content']}"
                    sub_para.font.size = Pt(14)
                    sub_para.font.color.rgb = PRIMARY
                    sub_para.font.bold = True
                    y_pos += 0.35
        
        # Add footer
        footer = slide.shapes.add_textbox(
            Inches(0.5), Inches(7.1),
            Inches(12), Inches(0.3)
        )
        footer_frame = footer.text_frame
        footer_para = footer_frame.paragraphs[0]
        footer_para.text = "ITDLUPanel v2 - Technical Deep Dive | Đại học Đà Lạt © 2026"
        footer_para.font.size = Pt(10)
        footer_para.font.color.rgb = GRAY
        footer_para.alignment = PP_ALIGN.CENTER
    
    prs.save(output_path)
    print(f"PPTX created: {output_path}")

if __name__ == '__main__':
    # Read markdown
    md_path = '/root/ITDLUPanel/docs/slides/02-technical-deep-dive.md'
    pptx_path = '/root/ITDLUPanel/docs/slides/02-technical-deep-dive.pptx'
    
    with open(md_path, 'r', encoding='utf-8') as f:
        md_content = f.read()
    
    # Parse and create presentation
    slides = parse_slides(md_content)
    create_presentation(slides, pptx_path)
    
    print(f"Created {len(slides)} slides")
